from __future__ import annotations

from copy import copy
from datetime import datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from .ics205b_models import ICS205BForm

FORM_REVISION = "6/15/2018"
ASSIGNMENT_ROWS_PER_EXCEL_PAGE = 21


def _local(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return timezone.localtime(value) if timezone.is_aware(value) else value


def _date(value: datetime | None) -> str:
    value = _local(value)
    return value.strftime("%m/%d/%Y") if value else ""


def _time(value: datetime | None) -> str:
    value = _local(value)
    return value.strftime("%H:%M") if value else ""


def _prepared_by(form: ICS205BForm) -> str:
    return " — ".join(filter(None, [form.prepared_by_name, form.prepared_by_position]))


class _NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.setFont("Helvetica", 8)
            self.drawCentredString(
                landscape(letter)[0] / 2,
                0.22 * inch,
                f"Page {self._pageNumber} of {page_count}",
            )
            super().showPage()
        super().save()


def _draw_pdf_header_footer(pdf: canvas.Canvas, doc, form: ICS205BForm):
    page_width, page_height = landscape(letter)
    left = 0.25 * inch
    right = page_width - 0.25 * inch
    top = page_height - 0.25 * inch
    thin = 0.5

    pdf.saveState()
    pdf.setStrokeColor(colors.black)
    pdf.setLineWidth(thin)
    pdf.setFont("Helvetica-Bold", 10)

    # Header: title plus incident/date/operational-period blocks.
    header_bottom = top - 0.78 * inch
    title_right = left + 2.55 * inch
    incident_right = title_right + 1.72 * inch
    prepared_right = incident_right + 1.55 * inch
    pdf.rect(left, header_bottom, right - left, top - header_bottom)
    for x in [title_right, incident_right, prepared_right]:
        pdf.line(x, header_bottom, x, top)

    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawCentredString(
        (left + title_right) / 2,
        top - 0.28 * inch,
        "INCIDENT INFORMATION",
    )
    pdf.drawCentredString(
        (left + title_right) / 2,
        top - 0.48 * inch,
        "MANAGEMENT PLAN (ICS 205b)",
    )

    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawCentredString((title_right + incident_right) / 2, top - 0.16 * inch, "1. Incident Name")
    pdf.setFont("Helvetica", 9)
    pdf.drawCentredString(
        (title_right + incident_right) / 2, top - 0.49 * inch, form.incident.name[:34]
    )

    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawCentredString(
        (incident_right + prepared_right) / 2, top - 0.16 * inch, "2. Date/Time Prepared"
    )
    pdf.setFont("Helvetica", 8)
    pdf.drawString(
        incident_right + 0.08 * inch, top - 0.40 * inch, f"Date: {_date(form.prepared_at)}"
    )
    pdf.drawString(
        incident_right + 0.08 * inch, top - 0.62 * inch, f"Time: {_time(form.prepared_at)}"
    )

    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawCentredString(
        (prepared_right + right) / 2, top - 0.16 * inch, "3. Operational Period Date/Time"
    )
    pdf.setFont("Helvetica", 8)
    pdf.drawString(
        prepared_right + 0.08 * inch,
        top - 0.40 * inch,
        f"From: {_date(form.operational_period.starts_at)} {_time(form.operational_period.starts_at)}",
    )
    pdf.drawString(
        prepared_right + 0.08 * inch,
        top - 0.62 * inch,
        f"To: {_date(form.operational_period.ends_at)} {_time(form.operational_period.ends_at)}",
    )

    pdf.setFillColor(colors.HexColor("#E7E6E6"))
    pdf.rect(left, header_bottom - 0.28 * inch, right - left, 0.28 * inch, stroke=1, fill=1)
    pdf.setFillColor(colors.black)
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawCentredString(
        page_width / 2,
        header_bottom - 0.19 * inch,
        "4. Information Technology Infrastructure & Services Assignment",
    )

    # Footer: repeat all prepared-by/location information on every page.
    footer_top = 0.90 * inch
    footer_bottom = 0.34 * inch
    pdf.rect(left, footer_bottom, right - left, footer_top - footer_bottom)
    footer_width = right - left
    cuts = [0.25, 0.37, 0.49, 0.62, 0.78, 0.86, 0.93]
    for ratio in cuts:
        x = left + footer_width * ratio
        pdf.line(x, footer_bottom, x, footer_top)
    pdf.setFont("Helvetica-Bold", 6.8)
    labels = [
        (left + 0.03 * inch, "5. Prepared By (Name and Position)"),
        (left + footer_width * 0.25 + 0.03 * inch, "Phone Number"),
        (left + footer_width * 0.37 + 0.03 * inch, "Signature"),
        (left + footer_width * 0.49 + 0.03 * inch, "Date / Time"),
        (left + footer_width * 0.62 + 0.03 * inch, "6. Incident Location"),
        (left + footer_width * 0.78 + 0.03 * inch, "State"),
        (left + footer_width * 0.86 + 0.03 * inch, "County"),
        (left + footer_width * 0.93 + 0.03 * inch, "City"),
    ]
    for x, text in labels:
        pdf.drawString(x, footer_top - 0.14 * inch, text)
    pdf.setFont("Helvetica", 6.8)
    values = [
        (left + 0.03 * inch, _prepared_by(form)),
        (left + footer_width * 0.25 + 0.03 * inch, form.prepared_by_phone),
        (left + footer_width * 0.37 + 0.03 * inch, form.prepared_by_signature),
        (
            left + footer_width * 0.49 + 0.03 * inch,
            f"{_date(form.prepared_at)} {_time(form.prepared_at)}",
        ),
        (left + footer_width * 0.62 + 0.03 * inch, form.incident_location),
        (left + footer_width * 0.78 + 0.03 * inch, form.state),
        (left + footer_width * 0.86 + 0.03 * inch, form.county),
        (left + footer_width * 0.93 + 0.03 * inch, form.city),
    ]
    for x, text in values:
        pdf.drawString(x, footer_bottom + 0.12 * inch, (text or "")[:34])

    pdf.setFont("Helvetica", 7)
    pdf.drawString(left, 0.20 * inch, "ICS Form 205b")
    pdf.drawRightString(right, 0.20 * inch, f"Form Revision: {FORM_REVISION}")
    pdf.restoreState()


def render_ics205b_pdf(form: ICS205BForm) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.25 * inch,
        rightMargin=0.25 * inch,
        topMargin=1.42 * inch,
        bottomMargin=1.05 * inch,
        title="ICS Form 205b",
        author="ICT Branch Toolkit",
    )
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "ICS205BBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=6.3,
        leading=7.2,
        spaceAfter=0,
        spaceBefore=0,
    )
    header_style = ParagraphStyle(
        "ICS205BHeader",
        parent=body_style,
        fontName="Helvetica-Bold",
        fontSize=6.2,
        leading=7,
        alignment=TA_CENTER,
    )
    headings = [
        "Assignment",
        "IT Resource Type",
        "Name of Application or Resource",
        "Usage or Description",
        "Platform",
        "Developer",
        "Login/Install",
        "Equipment Location<br/>Web Address, IP Address or SSID",
        "POC Information",
        "Remarks",
    ]
    rows = [[Paragraph(value, header_style) for value in headings]]
    for item in form.assignments.all():
        rows.append(
            [
                Paragraph(value or "", body_style)
                for value in [
                    item.assignment,
                    item.it_resource_type,
                    item.resource_name,
                    item.usage_description,
                    item.platform,
                    item.developer,
                    item.login_install,
                    item.equipment_location,
                    item.poc_information,
                    item.remarks,
                ]
            ]
        )
    if len(rows) == 1:
        rows.append([Paragraph("", body_style) for _ in headings])

    widths = [0.68, 0.78, 0.95, 0.95, 0.72, 0.72, 0.78, 1.38, 0.82, 1.26]
    table = Table(rows, colWidths=[value * inch for value in widths], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7E6E6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 1), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 3),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ]
        )
    )
    doc.build(
        [table],
        onFirstPage=lambda pdf, current_doc: _draw_pdf_header_footer(pdf, current_doc, form),
        onLaterPages=lambda pdf, current_doc: _draw_pdf_header_footer(pdf, current_doc, form),
        canvasmaker=_NumberedCanvas,
    )
    return buffer.getvalue()


def _border() -> Border:
    thin = Side(style="thin", color="000000")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _build_excel_page(ws, form: ICS205BForm, assignments, page_number: int, page_count: int):
    ws.sheet_view.showGridLines = False
    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 18.7265625
    ws.column_dimensions["K"].width = 37.7265625
    heights = {1: 25, 2: 20.15, 3: 20.15, 4: 25, 5: 55, 27: 20.15, 28: 36, 29: 16}
    for row in range(6, 27):
        heights[row] = 36
    for row, height in heights.items():
        ws.row_dimensions[row].height = height

    merges = [
        "A1:C3",
        "D1:E1",
        "D2:E3",
        "F1:G1",
        "H1:K1",
        "A4:K4",
        "H5:I5",
        "A27:C27",
        "A28:C28",
        "G27:H27",
        "G28:H28",
        "B29:J29",
    ] + [f"H{row}:I{row}" for row in range(6, 27)]
    for ref in merges:
        ws.merge_cells(ref)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    bold = Font(bold=True, size=10)
    title = Font(bold=True, size=15)
    gray = PatternFill("solid", fgColor="E7E6E6")

    for row in ws.iter_rows(min_row=1, max_row=29, min_col=1, max_col=11):
        for cell in row:
            cell.border = _border()
            cell.alignment = left
            cell.font = Font(size=9)

    ws["A1"] = "INCIDENT INFORMATION\nMANAGEMENT PLAN (ICS 205b)"
    ws["A1"].font = title
    ws["A1"].alignment = center
    ws["D1"] = "1. Incident Name"
    ws["D1"].font = bold
    ws["D1"].alignment = center
    ws["D2"] = form.incident.name
    ws["D2"].alignment = center
    ws["F1"] = "2. Date/Time Prepared"
    ws["F1"].font = bold
    ws["F1"].alignment = center
    ws["F2"] = "Date:"
    ws["F3"] = "Time:"
    ws["G2"] = _date(form.prepared_at)
    ws["G3"] = _time(form.prepared_at)
    ws["H1"] = "3. Operational Period Date/Time"
    ws["H1"].font = bold
    ws["H1"].alignment = center
    ws["H2"] = "Date From:"
    ws["H3"] = "Time From:"
    ws["I2"] = _date(form.operational_period.starts_at)
    ws["I3"] = _time(form.operational_period.starts_at)
    ws["J2"] = "Date To:"
    ws["J3"] = "Time To:"
    ws["K2"] = _date(form.operational_period.ends_at)
    ws["K3"] = _time(form.operational_period.ends_at)

    ws["A4"] = "4. Information Technology Infrastructure & Services Assignment"
    ws["A4"].font = Font(bold=True, size=11)
    ws["A4"].alignment = center
    for cell in ws[4]:
        cell.fill = gray

    headings = [
        "Assignment",
        "IT Resource Type",
        "Name of\nApplication or\nResource",
        "Usage or\nDescription",
        "Platform",
        "Developer",
        "Login/Install",
        "Equipment Location\nWeb Address, IP Address or SSID",
        "POC Information",
        "Remarks",
    ]
    target_columns = ["A", "B", "C", "D", "E", "F", "G", "H", "J", "K"]
    for col, value in zip(target_columns, headings, strict=True):
        ws[f"{col}5"] = value
        ws[f"{col}5"].font = bold
        ws[f"{col}5"].alignment = center
    for cell in ws[5]:
        cell.fill = gray

    for row_number, item in enumerate(assignments, start=6):
        values = [
            item.assignment,
            item.it_resource_type,
            item.resource_name,
            item.usage_description,
            item.platform,
            item.developer,
            item.login_install,
            item.equipment_location,
            item.poc_information,
            item.remarks,
        ]
        for col, value in zip(target_columns, values, strict=True):
            ws[f"{col}{row_number}"] = value
            ws[f"{col}{row_number}"].alignment = left

    ws["A27"] = "5. Prepared By (Name and Position)"
    ws["D27"] = "Phone Number"
    ws["E27"] = "Signature"
    ws["F27"] = "Date / Time"
    ws["G27"] = "6. Incident Location"
    ws["I27"] = "State"
    ws["J27"] = "County"
    ws["K27"] = "City"
    for ref in ["A27", "D27", "E27", "F27", "G27", "I27", "J27", "K27"]:
        ws[ref].font = Font(bold=True, size=8)
    ws["A28"] = _prepared_by(form)
    ws["D28"] = form.prepared_by_phone
    ws["E28"] = form.prepared_by_signature
    ws["F28"] = f"{_date(form.prepared_at)} {_time(form.prepared_at)}".strip()
    ws["G28"] = form.incident_location
    ws["I28"] = form.state
    ws["J28"] = form.county
    ws["K28"] = form.city

    ws["A29"] = "ICS Form 205b"
    ws["K29"] = f"Form Revision: {FORM_REVISION}"
    ws["K29"].alignment = Alignment(horizontal="right", vertical="center")
    for cell in ws[29]:
        cell.font = Font(size=8)

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.35
    ws.page_margins.header = 0
    ws.page_margins.footer = 0.18
    ws.print_area = "A1:K29"
    ws.oddFooter.center.text = f"Page {page_number} of {page_count}"
    ws.oddFooter.center.size = 8


def render_ics205b_xlsx(form: ICS205BForm) -> bytes:
    assignments = list(form.assignments.all())
    chunks = [
        assignments[index : index + ASSIGNMENT_ROWS_PER_EXCEL_PAGE]
        for index in range(0, len(assignments), ASSIGNMENT_ROWS_PER_EXCEL_PAGE)
    ] or [[]]
    workbook = Workbook()
    workbook.remove(workbook.active)
    page_count = len(chunks)
    for index, chunk in enumerate(chunks, start=1):
        title = "ICS 205B" if index == 1 else f"ICS 205B ({index})"
        sheet = workbook.create_sheet(title)
        _build_excel_page(sheet, form, chunk, index, page_count)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
