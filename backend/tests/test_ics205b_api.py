from io import BytesIO

import pdfplumber
import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.authtoken.models import Token

from apps.accounts.models import Role, UserRoleAssignment
from apps.audit.models import AuditEvent
from apps.extensions.ics205b_models import ICS205BAssignment, ICS205BForm
from apps.incidents.models import Incident, IncidentMembership, OperationalPeriod


def auth_header(user):
    token, _ = Token.objects.get_or_create(user=user)
    return {"HTTP_AUTHORIZATION": f"Token {token.key}"}


def setup_scenario():
    admin = get_user_model().objects.create_superuser(
        "ics205b-admin",
        "ics205b-admin@example.invalid",
        "safe-test-password",
    )
    incident = Incident.objects.create(
        name="Synthetic ITSL Exercise",
        incident_number="SYN-205B",
        created_by=admin,
    )
    IncidentMembership.objects.create(
        incident=incident,
        user=admin,
        role=Role.ADMINISTRATOR,
        assigned_by=admin,
    )
    period = OperationalPeriod.objects.create(
        incident=incident,
        name="Operational Period 1",
        starts_at="2026-09-11T12:00:00Z",
        ends_at="2026-09-12T00:00:00Z",
        created_by=admin,
    )
    return admin, incident, period


def create_form(client, admin, incident, period):
    response = client.post(
        "/api/ics205b-forms/",
        {
            "incident": str(incident.id),
            "operational_period": str(period.id),
            "prepared_at": "2026-09-11T14:30:00Z",
            "prepared_by_name": "Synthetic ITSL",
            "prepared_by_position": "ITSL",
            "prepared_by_phone": "555-0100",
            "prepared_by_signature": "Synthetic ITSL",
            "incident_location": "Synthetic Command Post",
            "state": "TX",
            "county": "Synthetic County",
            "city": "Synthetic City",
        },
        content_type="application/json",
        **auth_header(admin),
    )
    assert response.status_code == 201, response.content
    return ICS205BForm.objects.get(pk=response.json()["id"])


@pytest.mark.django_db
def test_ics205b_create_edit_and_assignment_order(client):
    admin, incident, period = setup_scenario()
    form = create_form(client, admin, incident, period)
    headers = auth_header(admin)

    first = client.post(
        "/api/ics205b-assignments/",
        {
            "form": str(form.id),
            "assignment": "Command",
            "it_resource_type": "Application",
            "resource_name": "Synthetic Messaging",
            "usage_description": "Synthetic incident messaging",
            "platform": "Web",
            "developer": "Synthetic Vendor",
            "login_install": "Use approved account; no credentials stored here",
            "equipment_location": "https://example.invalid",
            "poc_information": "Synthetic POC",
            "remarks": "Exercise only",
        },
        content_type="application/json",
        **headers,
    )
    assert first.status_code == 201, first.content
    assert first.json()["position"] == 1

    second = client.post(
        "/api/ics205b-assignments/",
        {"form": str(form.id), "assignment": "Planning", "resource_name": "Synthetic GIS"},
        content_type="application/json",
        **headers,
    )
    assert second.status_code == 201
    assert second.json()["position"] == 2

    updated = client.patch(
        f"/api/ics205b-assignments/{second.json()['id']}/",
        {"remarks": "Updated synthetic note"},
        content_type="application/json",
        **headers,
    )
    assert updated.status_code == 200
    assert updated.json()["remarks"] == "Updated synthetic note"

    deleted = client.delete(
        f"/api/ics205b-assignments/{first.json()['id']}/",
        **headers,
    )
    assert deleted.status_code == 204
    remaining = ICS205BAssignment.objects.get(pk=second.json()["id"])
    assert remaining.position == 1


@pytest.mark.django_db
def test_ics205b_rejects_operational_period_from_another_incident(client):
    admin, incident, _ = setup_scenario()
    other = Incident.objects.create(name="Other Synthetic Incident", created_by=admin)
    other_period = OperationalPeriod.objects.create(
        incident=other,
        name="Other OP",
        starts_at="2026-09-11T12:00:00Z",
        ends_at="2026-09-11T18:00:00Z",
        created_by=admin,
    )
    response = client.post(
        "/api/ics205b-forms/",
        {"incident": str(incident.id), "operational_period": str(other_period.id)},
        content_type="application/json",
        **auth_header(admin),
    )
    assert response.status_code == 400
    assert "operational_period" in response.json()


@pytest.mark.django_db
def test_ics205b_multipage_excel_repeats_form_identity_and_page_numbers(client):
    admin, incident, period = setup_scenario()
    form = create_form(client, admin, incident, period)
    for position in range(1, 44):
        ICS205BAssignment.objects.create(
            form=form,
            position=position,
            assignment=f"Assignment {position}",
            it_resource_type="Synthetic service",
            resource_name=f"Synthetic Resource {position}",
            remarks="Synthetic export test",
        )

    response = client.get(f"/api/ics205b-forms/{form.id}/xlsx/", **auth_header(admin))
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["ICS 205B", "ICS 205B (2)", "ICS 205B (3)"]
    for page_number, sheet in enumerate(workbook.worksheets, start=1):
        assert "ICS 205b" in sheet["A1"].value
        assert sheet["D2"].value == "Synthetic ITSL Exercise"
        assert sheet["A29"].value == "ICS Form 205b"
        assert sheet["K29"].value == "Form Revision: 6/15/2018"
        assert sheet.oddFooter.center.text == f"Page {page_number} of 3"
        assert sheet.print_area == "'ICS 205B'!$A$1:$K$29" if page_number == 1 else sheet.print_area

    event = AuditEvent.objects.filter(action="ics205b.xlsx_exported").latest("occurred_at")
    assert event.details["byte_size"] == len(response.content)


@pytest.mark.django_db
def test_ics205b_multipage_pdf_repeats_headers_footers_and_page_numbers(client):
    admin, incident, period = setup_scenario()
    form = create_form(client, admin, incident, period)
    for position in range(1, 70):
        ICS205BAssignment.objects.create(
            form=form,
            position=position,
            assignment=f"Assignment {position}",
            it_resource_type="Synthetic network service",
            resource_name=f"Synthetic Resource {position}",
            usage_description=("Long synthetic description for pagination and wrapping. " * 3),
            equipment_location="https://example.invalid/synthetic/path",
            remarks="Synthetic only",
        )

    response = client.get(f"/api/ics205b-forms/{form.id}/pdf/", **auth_header(admin))
    assert response.status_code == 200
    with pdfplumber.open(BytesIO(response.content)) as document:
        assert len(document.pages) > 1
        total = len(document.pages)
        for page_number, page in enumerate(document.pages, start=1):
            text = page.extract_text() or ""
            assert "INCIDENT INFORMATION" in text
            assert "Synthetic ITSL Exercise" in text
            assert "Information Technology Infrastructure & Services Assignment" in text
            assert "ICS Form 205b" in text
            assert "Form Revision: 6/15/2018" in text
            assert f"Page {page_number} of {total}" in text

    event = AuditEvent.objects.filter(action="ics205b.pdf_exported").latest("occurred_at")
    assert event.details["byte_size"] == len(response.content)


@pytest.mark.django_db
def test_ics205b_export_denied_for_read_only_role(client):
    admin, incident, period = setup_scenario()
    form = create_form(client, admin, incident, period)
    reader = get_user_model().objects.create_user(
        "ics205b-reader",
        "reader@example.invalid",
        "safe-test-password",
    )
    UserRoleAssignment.objects.create(user=reader, role=Role.READ_ONLY)
    IncidentMembership.objects.create(
        incident=incident,
        user=reader,
        role=Role.READ_ONLY,
        assigned_by=admin,
    )
    response = client.get(f"/api/ics205b-forms/{form.id}/pdf/", **auth_header(reader))
    assert response.status_code == 403
    assert not AuditEvent.objects.filter(action="ics205b.pdf_exported").exists()
